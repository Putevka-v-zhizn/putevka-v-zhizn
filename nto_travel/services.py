import hashlib
import re
from io import BytesIO
from zipfile import BadZipFile

from django.core.exceptions import ValidationError
from django.db import transaction
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from core.models import Notification, UserNotification
from config import BASE_URL
from scholar_form.models import UserInfo
from .forms import CIRCUMSTANCE_CHOICES, POSSIBLE_PAYER_CHOICES, PREPARATION_CHOICES, PREVIOUS_PAYER_CHOICES, exclusive
from .models import NtoApplication, NtoApplicationEvent, NtoApplicationProfile, NtoCampaign, NtoProfile, NtoProfileImport


def create_application(user, campaign):
    info = UserInfo.objects.filter(user=user).first()
    account = getattr(user, "telegram_account", None)
    defaults = {
        "last_name": (info.last_name if info else "") or user.last_name,
        "first_name": (info.first_name if info else "") or user.first_name,
        "middle_name": info.middle_name if info else "",
        "region": info.region if info else "",
        "city": info.city if info else "",
        "school_class": info.next_year_class_digit if info else None,
        "phone": info.phone if info else "",
        "telegram": account.username if account and account.username else "",
        "email": user.email,
    }
    application, _ = NtoApplication.objects.get_or_create(user=user, campaign=campaign, defaults=defaults)
    return application


def save_ranked_profiles(application, profiles):
    if len({p.pk for p in profiles}) != len(profiles):
        raise ValidationError("Профиль выбран повторно.")
    if any(p.campaign_id != application.campaign_id for p in profiles):
        raise ValidationError("Профиль относится к другому сезону.")
    existing = set(application.ranked_profiles.values_list("profile_id", flat=True))
    if any(not p.is_active and p.pk not in existing for p in profiles):
        raise ValidationError("Новый выбор неактивного профиля недоступен.")
    with transaction.atomic():
        application.ranked_profiles.all().delete()
        NtoApplicationProfile.objects.bulk_create([
            NtoApplicationProfile(application=application, profile=p, rank=i)
            for i, p in enumerate(profiles, 1)
        ])


def _valid_list(values, choices, special=()):
    if not isinstance(values, list) or not all(isinstance(value, str) for value in values):
        return False
    if any(value not in dict(choices) for value in values):
        return False
    try:
        exclusive(values, special)
    except Exception:
        return False
    return True


def validate_submission(app):
    app.full_clean()
    errors = {}
    def require(condition, key, message="Обязательный ответ."):
        if not condition:
            errors[key] = message

    for key in ("last_name", "first_name", "region", "city", "phone", "email", "parent_name", "parent_role", "parent_phone"):
        require(bool((getattr(app, key) or "").strip()), key)
    for key in ("phone", "parent_phone"):
        require(bool(re.fullmatch(r"\+[0-9]{10,15}", getattr(app, key) or "")), key, "Телефон должен содержать от 10 до 15 цифр.")
    require(app.school_class is not None and 1 <= app.school_class <= 11, "school_class")
    require(app.finalist_declared, "finalist_declared")
    require(app.parent_contact_declared, "parent_contact_declared")
    ranks = list(app.ranked_profiles.select_related("profile").order_by("rank"))
    require(bool(ranks) and [p.rank for p in ranks] == list(range(1, len(ranks) + 1)), "profiles", "Выберите профили и задайте приоритеты подряд с 1.")
    require(all(row.profile.campaign_id == app.campaign_id and row.profile.is_active for row in ranks), "profiles", "Выберите активные профили этого сезона.")
    require(app.previous_final is not None, "previous_final")
    if app.previous_final is True:
        require(_valid_list(app.previous_payers, PREVIOUS_PAYER_CHOICES, {"local", "unknown"}) and bool(app.previous_payers), "previous_payers")
    require(_valid_list(app.possible_payers, POSSIBLE_PAYER_CHOICES, {"not_needed", "unknown", "none"}) and bool(app.possible_payers), "possible_payers")
    support_values = dict(NtoApplication.SupportAnswer.choices)
    require(app.school_support in support_values, "school_support")
    require(app.parent_support in support_values, "parent_support")
    require(app.ticket_cost is not None and 0 <= app.ticket_cost <= 10000000, "ticket_cost")
    require(app.income_disclosure in dict(NtoApplication.IncomeDisclosure.choices), "income_disclosure")
    if app.income_disclosure == NtoApplication.IncomeDisclosure.AMOUNT:
        require(app.income_per_member is not None and app.income_per_member >= 0, "income_per_member")
        require(app.income_proof is not None, "income_proof")
    require(app.family_material_status in dict(UserInfo.FamilyMaterialStatus.choices), "family_material_status")
    require(_valid_list(app.circumstances, CIRCUMSTANCE_CHOICES, {"none", "declined"}) and bool(app.circumstances), "circumstances")
    if set(app.circumstances or []) - {"none", "declined"}:
        require(app.circumstance_proof is not None, "circumstance_proof")
    if "other" in (app.circumstances or []):
        require(bool(app.situation_details.strip()), "situation_details")
    require(len(app.situation_details or "") <= 1000, "situation_details", "Не более 1000 знаков.")
    for key in ("universities", "specializations"):
        values = getattr(app, key)
        require(isinstance(values, list) and 1 <= len(values) <= 5 and all(isinstance(v, str) and 0 < len(v.strip()) <= 300 for v in values), key)
    require(isinstance(app.achievements, list) and len(app.achievements) <= 5 and all(isinstance(v, str) and 0 < len(v.strip()) <= 300 for v in app.achievements), "achievements")
    require(bool(app.achievements) != bool(app.no_achievements), "achievements", "Укажите достижения или отметьте их отсутствие.")
    require(_valid_list(app.preparation_methods, PREPARATION_CHOICES, {"none"}) and bool(app.preparation_methods), "preparation_methods")
    for key, limit in (("olympiad_plans", 3000), ("preparation_details", 3000), ("nto_goal", 2000), ("admission_without_trip", 2000)):
        value = (getattr(app, key) or "").strip()
        require(bool(value) and len(value) <= limit, key, f"Обязательный ответ, не более {limit} знаков.")
    require(app.refused_olympiads is not None, "refused_olympiads")
    if app.refused_olympiads is True:
        require(bool(app.refusal_reason.strip()) and len(app.refusal_reason) <= 1000, "refusal_reason")
    require(app.has_admission_advantages is not None, "has_admission_advantages")
    if app.has_admission_advantages is True:
        require(bool(app.advantage_details.strip()) and len(app.advantage_details) <= 1000, "advantage_details")
    require(len(app.free_thoughts or "") <= 1000, "free_thoughts", "Не более 1000 знаков.")
    if errors:
        raise ValidationError(errors)
    return ranks


def _event(app, actor, action, old_status, comment=""):
    return NtoApplicationEvent.objects.create(
        application=app, actor=actor, action=action, old_status=old_status,
        new_status=app.status, comment=comment,
    )


def _notify(user, message, application, sender=None):
    url = BASE_URL.rstrip("/") + reverse("nto_travel:application", kwargs={"slug": application.campaign.slug})
    notice = Notification.objects.create(sender=sender, message=f"{message}\n{url}")
    UserNotification.objects.create(notification=notice, recipient=user)


@transaction.atomic
def submit_application(application_id, actor, *, consent_confirmed):
    app = NtoApplication.objects.select_for_update().select_related("campaign").get(pk=application_id, user=actor)
    now = timezone.now()
    if app.status == NtoApplication.Status.DRAFT:
        if not app.campaign.accepting:
            raise ValidationError("Приём заявок сейчас закрыт.")
    elif app.status == NtoApplication.Status.CLARIFICATION:
        if not app.clarification_due_at or now >= app.clarification_due_at:
            raise ValidationError("Срок уточнения завершён.")
    else:
        raise ValidationError("Эта заявка уже отправлена.")
    if not consent_confirmed or not app.campaign.consent_revision_id:
        raise ValidationError("Необходимо принять согласие на обработку данных.")
    ranks = validate_submission(app)
    for row in ranks:
        row.name_snapshot = row.profile.name
        row.city_snapshot = row.profile.final_city
        row.save(update_fields=["name_snapshot", "city_snapshot"])
    old_status = app.status
    app.status = NtoApplication.Status.SUBMITTED
    app.final_city = ranks[0].city_snapshot
    app.email = actor.email
    app.last_submitted_at = now
    app.consent_revision = app.campaign.consent_revision
    app.consent_accepted_at = now
    app.save()
    _event(app, actor, "resubmitted" if old_status == NtoApplication.Status.CLARIFICATION else "submitted", old_status)
    _notify(actor, "Заявка на поддержку поездки НТО отправлена.", app)
    return app


@transaction.atomic
def review_application(application_id, actor, action, *, comment="", due_at=None):
    app = NtoApplication.objects.select_for_update().get(pk=application_id)
    if app.status != NtoApplication.Status.SUBMITTED:
        raise ValidationError("Рассматривать можно только поданную заявку.")
    old_status = app.status
    comment = (comment or "").strip()
    if action == "clarify":
        if not comment or not due_at or due_at <= timezone.now():
            raise ValidationError("Укажите комментарий и будущий срок уточнения.")
        app.status = NtoApplication.Status.CLARIFICATION
        app.clarification_due_at = due_at
        app.public_comment = comment
        message = f"Заявка НТО возвращена на уточнение до {timezone.localtime(due_at):%d.%m.%Y %H:%M}. Комментарий: {comment}"
    elif action == "approve":
        app.status = NtoApplication.Status.APPROVED
        app.decided_by = actor
        app.decided_at = timezone.now()
        app.public_comment = comment
        message = "Заявка на поддержку поездки НТО одобрена. Подробности в личном кабинете."
    elif action == "reject":
        if not comment:
            raise ValidationError("Укажите причину отказа.")
        app.status = NtoApplication.Status.REJECTED
        app.decided_by = actor
        app.decided_at = timezone.now()
        app.public_comment = comment
        message = "По заявке на поддержку поездки НТО принято решение. Подробности в личном кабинете."
    else:
        raise ValidationError("Неизвестное действие.")
    app.save()
    _event(app, actor, action, old_status, comment)
    _notify(app.user, message, app, sender=actor)
    return app


def parse_profile_workbook(upload):
    data = upload.read()
    try:
        book = load_workbook(BytesIO(data), read_only=True, data_only=True)
        rows = book.active.iter_rows(values_only=True)
        header = [str(v or "").strip() for v in next(rows)]
        if header[:3] != ["external_code", "name", "final_city"]:
            raise ValueError("Ожидаются столбцы external_code, name, final_city.")
        parsed, errors, seen = [], [], {}
        for number, values in enumerate(rows, 2):
            if not any(values):
                continue
            if number > 5002:
                errors.append("Не более 5000 строк.")
                break
            code, name, city = [str(values[i] or "").strip() if i < len(values) else "" for i in range(3)]
            if not code or not name or not city or len(code) > 100 or len(name) > 255 or len(city) > 255:
                errors.append(f"Строка {number}: заполните код, название и город в допустимой длине.")
            elif code in seen:
                errors.append(f"Строка {number}: повтор кода {code}; прежний город: {seen[code]}.")
            else:
                seen[code] = city
                parsed.append({"external_code": code, "name": name, "final_city": city})
        if not parsed:
            errors.append("В файле нет профилей.")
        return parsed, errors, hashlib.sha256(data).hexdigest()
    except (ValueError, StopIteration, OSError, KeyError, BadZipFile, InvalidFileException) as exc:
        return [], [str(exc)], hashlib.sha256(data).hexdigest()


@transaction.atomic
def apply_profile_import(batch_id, campaign, actor):
    NtoCampaign.objects.select_for_update().get(pk=campaign.pk)
    batch = NtoProfileImport.objects.select_for_update().get(pk=batch_id, campaign=campaign, actor=actor)
    if batch.applied_at:
        raise ValidationError("Этот импорт уже подтверждён.")
    codes = []
    for row in batch.rows:
        codes.append(row["external_code"])
        NtoProfile.objects.update_or_create(
            campaign=campaign, external_code=row["external_code"],
            defaults={"name": row["name"], "final_city": row["final_city"], "is_active": True},
        )
    NtoProfile.objects.filter(campaign=campaign).exclude(external_code__in=codes).update(is_active=False)
    batch.applied_at = timezone.now()
    batch.save(update_fields=["applied_at"])
    return batch
