from datetime import timedelta
import json
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from core.models import RegistrationPersonalData, UserNotification, UserProgramPreferences
from scholar_form.models import UserInfo
from .models import (
    NtoApplication, NtoApplicationEvent, NtoApplicationProfile,
    NtoCampaign, NtoConsentRevision, NtoProfile, NtoProfileImport,
)
from .forms import FinanceForm, MotivationForm, PersonalForm, ProfileRankingForm
from .validators import normalize_phone
from core.services.programs import select_program
from .services import (
    apply_profile_import, create_application, parse_profile_workbook,
    review_application, submit_application, validate_submission,
    save_ranked_profiles,
)


class RussianPhoneTests(SimpleTestCase):
    def test_russian_formats_and_old_plus_eight_are_normalized(self):
        for raw in ('+7 (903) 123-45-67', '8 (903) 123-45-67', '79031234567', '9031234567', '+89031234567'):
            with self.subTest(raw=raw):
                self.assertEqual(normalize_phone(raw), '+79031234567')
        self.assertEqual(normalize_phone('8 (495) 123-45-67'), '+74951234567')
        self.assertEqual(normalize_phone('8 (800) 123-45-67'), '+78001234567')
        self.assertEqual(normalize_phone(''), '')

    def test_invalid_and_non_russian_numbers_are_rejected(self):
        for raw in ('+12025550123', '+77011234567', '1234567890', '903123456', '890312345678', 'abc89031234567', '++79031234567', '79031234567 доб. 123'):
            with self.subTest(raw=raw), self.assertRaises(ValidationError):
                normalize_phone(raw)


class NtoTravelTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="nto@example.test", email="nto@example.test", password="testpass123")
        cls.other = get_user_model().objects.create_user(username="other@example.test", email="other@example.test", password="testpass123")
        cls.staff = get_user_model().objects.create_user(username="staff@example.test", email="staff@example.test", password="testpass123", is_staff=True)
        cls.consent = NtoConsentRevision.objects.create(version="test-1", text="Тестовый текст согласия")
        now = timezone.now()
        cls.campaign = NtoCampaign.objects.create(
            slug="test-season", academic_year="2099/2100", name="Сезон НТО",
            opens_at=now - timedelta(days=1), closes_at=now + timedelta(days=1),
            state=NtoCampaign.State.DRAFT, consent_revision=cls.consent,
        )
        cls.profile = NtoProfile.objects.create(campaign=cls.campaign, external_code="robot", name="Робототехника", final_city="Томск")
        cls.campaign.state = NtoCampaign.State.PUBLISHED
        cls.campaign.save(update_fields=["state"])

    def test_profile_export_permissions_and_roundtrip(self):
        url = reverse("nto_staff:profile_export", args=[self.campaign.slug])
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.staff.groups.add(Group.objects.get(name="Проверяющий НТО"))
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.staff.groups.add(Group.objects.get(name="Администратор НТО"))
        NtoProfile.objects.create(campaign=self.campaign, external_code="old", name="Старый", final_city="Москва", is_active=False)
        self.profile.name = "=текст"
        self.profile.save()
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['B2'].data_type, 's')
        rows, errors, _ = parse_profile_workbook(BytesIO(response.content))
        self.assertEqual(errors, [])
        self.assertEqual(rows, [{"external_code": "robot", "name": "=текст", "final_city": "Томск"}])

    def test_history_separates_access_events_and_escapes_comments(self):
        app = create_application(self.user, self.campaign)
        self.staff.groups.add(Group.objects.get(name="Проверяющий НТО"))
        self.staff.first_name = "Анна"
        self.staff.last_name = "Иванова"
        self.staff.save()
        self.client.force_login(self.staff)
        first = NtoApplicationEvent.objects.create(application=app, action="submitted", actor=self.user, old_status="draft", new_status="submitted")
        second = NtoApplicationEvent.objects.create(application=app, action="clarify", actor=self.staff, old_status="submitted", new_status="clarification", comment="Первая строка\n<script>alert(1)</script>")
        NtoApplicationEvent.objects.filter(pk__in=[first.pk, second.pk]).update(created_at=first.created_at)
        NtoApplicationEvent.objects.bulk_create([
            NtoApplicationEvent(application=app, action="staff_view", actor=self.staff)
            for _ in range(55)
        ])
        NtoApplicationEvent.objects.create(application=app, action="export")
        response = self.client.get(reverse("nto_staff:detail", args=[self.campaign.slug, app.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual([row["label"] for row in response.context["events"]], ["Запрошено уточнение", "Заявка отправлена"])
        self.assertEqual(response.context["access_count"], 57)
        self.assertEqual(len(response.context["access_events"]), 50)
        self.assertContains(response, "Анна Иванова")
        self.assertContains(response, self.user.username)
        self.assertContains(response, "Система")
        self.assertContains(response, "Подана → На уточнении")
        self.assertContains(response, "Первая строка\n&lt;script&gt;alert(1)&lt;/script&gt;")
        self.assertNotContains(response, "<script>alert(1)</script>")
        self.assertNotContains(response, "<details open")
        self.assertContains(response, "Показаны последние 50 из 57 событий.")

    def test_all_event_actions_have_readable_labels(self):
        from .presentation import EVENT_LABELS, event_rows
        events = [NtoApplicationEvent(action=action) for action in EVENT_LABELS]
        self.assertEqual([row["label"] for row in event_rows(events)], list(EVENT_LABELS.values()))
        self.assertTrue(all(row["actor"] == "Система" and not row["transition"] for row in event_rows(events)))

    def test_export_values_order_and_empty_workbook(self):
        self.staff.groups.add(Group.objects.get(name="Администратор НТО"))
        self.client.force_login(self.staff)
        app = create_application(self.user, self.campaign)
        app.last_submitted_at = timezone.now()
        app.status = "submitted"
        app.ticket_cost = 0
        app.income_disclosure = "declined"
        app.school_support = "partial"
        app.possible_payers = ["region", "parents"]
        app.phone = "+79031234567"
        app.save()
        NtoApplicationProfile.objects.create(application=app, profile=self.profile, rank=1, name_snapshot="Сохранённый профиль", city_snapshot="Казань")
        draft = create_application(self.other, self.campaign)
        url = reverse("nto_staff:export", args=[self.campaign.slug])
        book = load_workbook(BytesIO(self.client.get(url).content))
        for sheet in book:
            self.assertEqual([sheet.cell(i, 1).value for i in (2, 3)], [app.pk, draft.pk])
            self.assertEqual(sheet.freeze_panes, "C2")
            self.assertEqual(len(sheet.tables), 1)
            self.assertIsNone(sheet.auto_filter.ref)
            table = next(iter(sheet.tables.values()))
            self.assertEqual(table.autoFilter.ref, table.ref)
            self.assertEqual(len({column.name.casefold() for column in table.tableColumns}), sheet.max_column)
        registry = dict(zip([c.value for c in book['Реестр'][1]], book['Реестр'][2]))
        self.assertEqual(registry['Стоимость билетов, ₽'].value, 0)
        self.assertEqual(registry['Стоимость билетов, ₽'].data_type, 'n')
        self.assertEqual(registry['Дата последней отправки'].data_type, 'd')
        self.assertIsNone(registry['Доход на человека, ₽/мес'].value)
        self.assertNotIn('Готовность указать доход', registry)
        self.assertEqual(registry['Ответ школы'].value, 'Оплатят частично')
        self.assertEqual(registry['Возможные источники оплаты'].value, 'Регион\nРодители')
        self.assertEqual(registry['Первый профиль'].value, 'Сохранённый профиль')
        self.assertEqual(registry['Город финала'].value, 'Казань')
        answers = dict(zip([c.value for c in book['Ответы анкеты'][1]], book['Ответы анкеты'][2]))
        self.assertEqual(answers['Телефон'].value, app.phone)
        self.assertEqual(answers['Телефон'].data_type, 's')
        self.assertNotIn('ID пользователя', answers)
        self.assertNotIn('Редакция согласия', answers)
        self.assertEqual(NtoApplicationEvent.objects.filter(application=app, action='export').count(), 1)
        empty = load_workbook(BytesIO(self.client.get(url, {'region': 'Несуществующий регион'}).content))
        self.assertEqual(empty.sheetnames, ['Реестр', 'Ответы анкеты'])
        for sheet in empty:
            self.assertEqual(sheet.max_row, 1)
            self.assertTrue(sheet.auto_filter.ref)

    def test_staff_edit_preserves_status_and_records_changes(self):
        app = self.filled_application()
        app.status = "approved"
        app.save()
        url = reverse('nto_staff:edit', args=[self.campaign.slug, app.pk])
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.staff.groups.add(Group.objects.get(name='Проверяющий НТО'))
        self.client.force_login(self.staff)
        self.assertEqual(self.client.post(url, {}).status_code, 403)
        self.staff.groups.add(Group.objects.get(name='Администратор НТО'))
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'mc.yandex.ru')
        payload = {}
        for form in response.context['forms']:
            for name in form.fields:
                value = form[name].value()
                if value is not None:
                    payload[form.add_prefix(name)] = value
        payload.update({'profile_order': json.dumps([self.profile.pk]), 'personal-first_name': 'Новое имя', 'finance-ticket_cost': '12345'})
        response = self.client.post(url, payload)
        self.assertEqual(response.status_code, 302)
        app.refresh_from_db()
        self.assertEqual(app.first_name, 'Новое имя')
        self.assertEqual(app.ticket_cost, 12345)
        self.assertEqual(app.status, 'approved')
        self.assertTrue(app.finalist_declared)
        self.assertTrue(app.parent_contact_declared)
        self.assertEqual(app.events.filter(action='staff_edit').count(), 1)
        comment = app.events.get(action='staff_edit').comment
        self.assertIn('Имя', comment)
        self.assertIn('Стоимость билетов, ₽', comment)
        self.assertNotIn('Ранее проходил в финал', comment)
        self.assertNotIn('Приоритетные вузы', comment)
        self.assertNotIn('Готовность подтвердить доход', comment)
        self.assertEqual(self.client.post(url, payload).status_code, 302)
        self.assertEqual(app.events.filter(action='staff_edit').count(), 1)
        payload['finance-ticket_cost'] = '-1'
        self.assertEqual(self.client.post(url, payload).status_code, 200)
        app.refresh_from_db()
        self.assertEqual(app.ticket_cost, 12345)
        self.assertEqual(app.events.filter(action='staff_edit').count(), 1)

    def test_word_questionnaire_fields_and_adult_condition(self):
        personal = PersonalForm()
        self.assertNotIn('finalist_declared', personal.fields)
        self.assertNotIn('parent_contact_declared', personal.fields)
        self.assertNotIn('parent_role', personal.fields)
        finance = FinanceForm()
        self.assertNotIn('income_disclosure', finance.fields)
        self.assertNotIn('not_discussed', dict(finance.fields['school_support'].choices))
        self.assertNotIn('declined', dict(finance.fields['circumstances'].choices))
        self.assertEqual(len(list(finance.fields['family_material_status'].choices)), 5)
        motivation = MotivationForm()
        self.assertNotIn('no_achievements', motivation.fields)
        self.assertNotIn('none', dict(motivation.fields['preparation_methods'].choices))
        app = self.filled_application()
        app.is_adult = True
        app.parent_name = app.parent_phone = ''
        app.finalist_declared = app.parent_contact_declared = False
        app.nto_goal = 'Цель ' * 1000
        app.achievements = ['Достижение ' * 100]
        app.preparation_methods = []
        validate_submission(app)
        app.is_adult = False
        with self.assertRaises(ValidationError) as raised:
            validate_submission(app)
        self.assertIn('parent_name', raised.exception.message_dict)
        self.assertIn('parent_phone', raised.exception.message_dict)

    def test_participant_fields_are_disabled_when_saving_is_unavailable(self):
        select_program(self.user, 'nto')
        self.client.force_login(self.user)
        app = self.filled_application()
        for status, due in [('submitted', None), ('approved', None), ('rejected', None), ('clarification', timezone.now() - timedelta(minutes=1))]:
            app.status, app.clarification_due_at = status, due
            app.save()
            for step in ('personal', 'finance', 'motivation'):
                with self.subTest(status=status, step=step):
                    url = reverse('nto_travel:apply_step', args=[self.campaign.slug, step])
                    response = self.client.get(url)
                    self.assertFalse(response.context['editable'])
                    self.assertTrue(all(field.disabled for field in response.context['form'].fields.values()))
                    self.assertNotContains(response, 'Сохранить и продолжить')
                    self.assertEqual(self.client.post(url, {'first_name': 'Изменение'}).status_code, 302)
            app.refresh_from_db()
            self.assertEqual(app.first_name, 'Иван')
        app.status = 'draft'
        app.save()
        self.campaign.closes_at = timezone.now() - timedelta(seconds=1)
        self.campaign.save()
        response = self.client.get(reverse('nto_travel:apply_step', args=[self.campaign.slug, 'finance']))
        self.assertTrue(all(field.disabled for field in response.context['form'].fields.values()))
        app.status = 'clarification'
        app.clarification_due_at = timezone.now() + timedelta(days=1)
        app.save()
        response = self.client.get(reverse('nto_travel:apply_step', args=[self.campaign.slug, 'finance']))
        self.assertTrue(response.context['editable'])
        self.assertTrue(all(not field.disabled for field in response.context['form'].fields.values()))

    def filled_application(self):
        app = create_application(self.user, self.campaign)
        for key, value in {
            "last_name": "Иванов", "first_name": "Иван", "region": "Томская область",
            "city": "Северск", "school_class": 10, "phone": "+79000000000",
            "email": self.user.email, "telegram": "example_user", "parent_name": "Иванова Мария", "parent_role": "мать",
            "parent_phone": "+79000000001", "finalist_declared": True,
            "parent_contact_declared": True, "previous_final": False,
            "previous_payers": [], "possible_payers": ["parents"],
            "school_support": "full", "parent_support": "partial",
            "ticket_cost": 1000, "income_per_member": 25000, "income_proof": True, "circumstance_proof": True, "is_adult": False,
            "family_material_status": UserInfo.FamilyMaterialStatus.values[0],
            "circumstances": ["none"], "universities": ["ТГУ"],
            "specializations": ["Информатика"], "no_achievements": True,
            "olympiad_plans": "Планирую НТО", "preparation_methods": ["self_study"],
            "preparation_details": "Занимаюсь еженедельно", "nto_goal": "Участие в финале",
            "refused_olympiads": False, "admission_without_trip": "Буду искать другие пути",
            "has_admission_advantages": False,
        }.items():
            setattr(app, key, value)
        app.save()
        NtoApplicationProfile.objects.get_or_create(application=app, profile=self.profile, defaults={"rank": 1})
        return app

    def test_unique_draft_and_submission_snapshot_notification(self):
        app = self.filled_application()
        self.assertEqual(create_application(self.user, self.campaign).pk, app.pk)
        submitted = submit_application(app.pk, self.user, consent_confirmed=True)
        self.assertEqual(submitted.status, NtoApplication.Status.SUBMITTED)
        self.assertEqual(submitted.final_city, "Томск")
        self.assertEqual(submitted.ranked_profiles.get().city_snapshot, "Томск")
        self.assertEqual(UserNotification.objects.filter(recipient=self.user).count(), 1)
        with self.assertRaises(ValidationError):
            submit_application(app.pk, self.user, consent_confirmed=True)
        self.assertEqual(NtoApplicationEvent.objects.filter(application=app, action="submitted").count(), 1)

    def test_both_phone_fields_and_submission_use_russian_validation(self):
        app = self.filled_application()
        form = PersonalForm(data={'phone': '8 (903) 123-45-67', 'parent_phone': '4951234567'}, instance=app)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data['phone'], '+79031234567')
        self.assertEqual(form.cleaned_data['parent_phone'], '+74951234567')
        form = PersonalForm(data={'phone': '+12025550123', 'parent_phone': 'abc'}, instance=app)
        self.assertFalse(form.is_valid())
        self.assertIn('phone', form.errors)
        self.assertIn('parent_phone', form.errors)
        app.refresh_from_db()
        app.phone = '89031234567'
        app.parent_phone = '+89031234567'
        validate_submission(app)
        self.assertEqual(app.phone, '+79031234567')
        self.assertEqual(app.parent_phone, '+79031234567')
        app.parent_phone = '+77011234567'
        with self.assertRaises(ValidationError) as raised:
            validate_submission(app)
        self.assertIn('parent_phone', raised.exception.message_dict)

    def test_conditional_answers_and_closed_campaign(self):
        app = self.filled_application()
        app.ticket_cost = 10000001
        with self.assertRaises(ValidationError):
            validate_submission(app)
        app.ticket_cost = 1000
        app.nto_goal = "x" * 2001
        validate_submission(app)
        app.nto_goal = "Участие в финале"
        app.possible_payers = ["none", "parents"]
        app.save()
        with self.assertRaises(ValidationError):
            validate_submission(app)
        app.possible_payers = ["none"]
        app.circumstances = ["other"]
        app.save()
        validate_submission(app)
        app.circumstance_proof = None
        app.situation_details = "Стоимость поездки недоступна семье"
        app.save()
        with self.assertRaises(ValidationError):
            validate_submission(app)  # Не указан ответ о подтверждении обстоятельств.
        app.circumstance_proof = False
        app.save()
        validate_submission(app)
        self.campaign.closes_at = timezone.now() - timedelta(seconds=1)
        self.campaign.save()
        with self.assertRaises(ValidationError):
            submit_application(app.pk, self.user, consent_confirmed=True)

    def test_review_is_final_and_clarification_can_outlive_campaign(self):
        app = self.filled_application()
        submit_application(app.pk, self.user, consent_confirmed=True)
        due = timezone.now() + timedelta(days=3)
        review_application(app.pk, self.staff, "clarify", comment="Уточните поездку", due_at=due)
        self.campaign.closes_at = timezone.now() - timedelta(seconds=1)
        self.campaign.save()
        app = submit_application(app.pk, self.user, consent_confirmed=True)
        self.assertEqual(app.status, NtoApplication.Status.SUBMITTED)
        app = review_application(app.pk, self.staff, "reject", comment="Недостаточно средств программы")
        self.assertEqual(app.status, NtoApplication.Status.REJECTED)
        with self.assertRaises(ValidationError):
            review_application(app.pk, self.staff, "approve")
        self.assertEqual(UserNotification.objects.filter(recipient=self.user).count(), 4)

    def test_import_preview_and_repeat_keeps_snapshots(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["external_code", "name", "final_city"])
        sheet.append(["robot", "Робототехника", "Москва"])
        sheet.append(["robot", "Робототехника", "Казань"])
        stream = BytesIO()
        book.save(stream)
        _, errors, _ = parse_profile_workbook(BytesIO(stream.getvalue()))
        self.assertTrue(errors)
        app = self.filled_application()
        submit_application(app.pk, self.user, consent_confirmed=True)
        batch = NtoProfileImport.objects.create(
            campaign=self.campaign, actor=self.staff, file_sha256="x" * 64,
            row_count=1, rows=[{"external_code": "robot", "name": "Роботы", "final_city": "Москва"}],
        )
        apply_profile_import(batch.pk, self.campaign, self.staff)
        self.assertEqual(NtoProfile.objects.get(pk=self.profile.pk).final_city, "Москва")
        self.assertEqual(app.ranked_profiles.get().city_snapshot, "Томск")

    def test_private_pages_permissions_and_xlsx(self):
        app = self.filled_application()
        app.last_name = "=FORMULA()"
        app.save(update_fields=["last_name"])
        self.client.force_login(self.user)
        response = self.client.get(reverse("nto_travel:review", args=[self.campaign.slug]))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "mc.yandex.ru")
        self.assertEqual(self.client.get(reverse("nto_staff:queue")).status_code, 403)
        self.client.force_login(self.staff)
        reviewer = Group.objects.get(name="Проверяющий НТО")
        self.staff.groups.add(reviewer)
        self.assertEqual(self.client.get(reverse("nto_staff:queue")).status_code, 200)
        detail = self.client.get(reverse("nto_staff:detail", args=[self.campaign.slug, app.pk]))
        self.assertEqual(detail.status_code, 200)
        self.assertNotContains(detail, "mc.yandex.ru")
        self.assertEqual(NtoApplicationEvent.objects.filter(application=app, action="staff_view").count(), 1)
        self.assertEqual(self.client.get(reverse("nto_staff:export", args=[self.campaign.slug])).status_code, 403)
        administrator = Group.objects.get(name="Администратор НТО")
        self.staff.groups.add(administrator)
        self.staff = get_user_model().objects.get(pk=self.staff.pk)
        template = self.client.get(reverse("nto_staff:profile_template", args=[self.campaign.slug]))
        self.assertEqual(load_workbook(BytesIO(template.content)).active["A1"].value, "external_code")
        response = self.client.get(reverse("nto_staff:export", args=[self.campaign.slug]))
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book.sheetnames, ["Реестр", "Ответы анкеты"])
        formula_cell = next(cell for row in book['Ответы анкеты'] for cell in row if cell.value == '=FORMULA()')
        self.assertEqual(formula_cell.data_type, 's')
        self.assertEqual(sheet.freeze_panes, 'C2')
        self.assertEqual(len(sheet.tables), 1)
        self.assertEqual(NtoApplicationEvent.objects.filter(application=app, action="export").count(), 1)
        filtered = self.client.get(reverse("nto_staff:export", args=[self.campaign.slug]), {"status": "submitted"})
        self.assertEqual(load_workbook(BytesIO(filtered.content)).active.max_row, 1)

    def test_incomplete_draft_and_conditional_form_cleanup(self):
        select_program(self.user, 'nto')
        self.client.force_login(self.user)
        url = reverse("nto_travel:apply_step", args=[self.campaign.slug, "personal"])
        response = self.client.post(url, {
            "first_name": "Иван", "profile_order": json.dumps([self.profile.pk]),
        })
        self.assertEqual(response.status_code, 302)
        app = NtoApplication.objects.get(user=self.user, campaign=self.campaign)
        self.assertEqual(app.status, NtoApplication.Status.DRAFT)
        self.assertEqual(app.ranked_profiles.count(), 1)
        self.assertEqual(app.email, self.user.email)
        response = self.client.post(reverse("nto_travel:apply_step", args=[self.campaign.slug, "finance"]), {
            "previous_final": "no", "previous_payers": ["region"],
            "income_disclosure": "declined", "income_per_member": "1000", "income_proof": "yes",
            "circumstances": ["none"], "circumstance_proof": "yes",
        })
        self.assertEqual(response.status_code, 302)
        app.refresh_from_db()
        self.assertEqual(app.previous_payers, [])
        self.assertEqual(app.income_per_member, 1000)
        self.assertTrue(app.income_proof)
        self.assertTrue(app.circumstance_proof)
        self.assertEqual(FinanceForm(instance=app)["previous_final"].value(), "no")
        self.assertEqual(MotivationForm(instance=app)["refused_olympiads"].value(), "")

    def test_income_amount_survives_save_and_reload_without_explicit_disclosure(self):
        select_program(self.user, 'nto')
        self.client.force_login(self.user)
        url = reverse("nto_travel:apply_step", args=[self.campaign.slug, "finance"])
        for disclosure in ("", "amount"):
            with self.subTest(disclosure=disclosure):
                response = self.client.post(url, {
                    "income_disclosure": disclosure, "income_per_member": "25\u00a0400,00",
                    "income_proof": "yes", "ticket_cost": "20 347",
                })
                self.assertEqual(response.status_code, 302)
                app = NtoApplication.objects.get(user=self.user, campaign=self.campaign)
                self.assertEqual(app.income_per_member, 25400)
                self.assertEqual(app.ticket_cost, 20347)
                self.assertNotIn("income_disclosure", FinanceForm().fields)
                self.assertIs(app.income_proof, True)
                response = self.client.get(url)
                self.assertEqual(response.context["form"]["income_per_member"].value(), 25400)
                self.assertContains(response, 'type="text" name="income_per_member"')
                self.assertContains(response, 'type="text" name="ticket_cost"')

    def test_picker_order_schema_and_available_profiles(self):
        app = create_application(self.user, self.campaign)
        second = NtoProfile.objects.create(campaign=self.campaign, external_code='data', name='Анализ данных', final_city='Москва')
        inactive = NtoProfile.objects.create(campaign=self.campaign, external_code='old', name='Старый профиль', final_city='Казань', is_active=False)
        other_campaign = NtoCampaign.objects.create(slug='other', academic_year='other', name='Другой сезон', opens_at=timezone.now(), closes_at=timezone.now()+timedelta(days=1))
        foreign = NtoProfile.objects.create(campaign=other_campaign, external_code='foreign', name='Другой профиль', final_city='Пермь')
        for payload in ([self.profile.pk, self.profile.pk], [foreign.pk], [inactive.pk], [True], [1.5], {'id': self.profile.pk}, 'bad'):
            with self.subTest(payload=payload):
                form = ProfileRankingForm({'profile_order': json.dumps(payload)}, campaign=self.campaign, application=app)
                self.assertFalse(form.is_valid())
        form = ProfileRankingForm({'profile_order': json.dumps([second.pk, self.profile.pk])}, campaign=self.campaign, application=app)
        self.assertTrue(form.is_valid(), form.errors)
        save_ranked_profiles(app, form.selected_profiles)
        self.assertEqual(list(app.ranked_profiles.values_list('profile_id', 'rank')), [(second.pk, 1), (self.profile.pk, 2)])
        self.assertEqual(ProfileRankingForm(campaign=self.campaign, application=app).selected_ids, [str(second.pk), str(self.profile.pk)])
        empty = ProfileRankingForm({'profile_order': '[]'}, campaign=self.campaign, application=app)
        self.assertTrue(empty.is_valid(), empty.errors)
        self.assertEqual(empty.selected_profiles, [])
        stale_page = ProfileRankingForm({f'choose_{self.profile.pk}': 'on', f'rank_{self.profile.pk}': '1'}, campaign=self.campaign, application=app)
        self.assertFalse(stale_page.is_valid())

    def test_picker_preserves_order_when_another_field_is_invalid(self):
        app = self.filled_application()
        second = NtoProfile.objects.create(campaign=self.campaign, external_code='data', name='Анализ данных', final_city='Москва')
        self.client.force_login(self.user)
        response = self.client.post(reverse('nto_travel:apply_step', args=[self.campaign.slug, 'personal']), {
            'phone': 'invalid', 'profile_order': json.dumps([second.pk, self.profile.pk]),
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['ranking'].selected_ids, [str(second.pk), str(self.profile.pk)])
        self.assertContains(response, 'nto-profile-search')
        self.assertNotContains(response, 'id_rank_')
        self.assertEqual(list(app.ranked_profiles.values_list('profile_id', flat=True)), [self.profile.pk])

    def test_other_user_cannot_read_application(self):
        self.filled_application()
        select_program(self.other, 'nto')
        self.client.force_login(self.other)
        response = self.client.get(reverse("nto_travel:application", args=[self.campaign.slug]))
        self.assertEqual(response.status_code, 404)

    def test_program_navigation_is_separate(self):
        info = UserInfo.objects.create(user=self.user)
        self.client.force_login(self.user)
        response = self.client.get(reverse("index"))
        self.assertContains(response, 'id="selectionMenu"')
        self.assertContains(response, 'id="selectionDropdown"')
        self.assertNotContains(response, 'href="' + reverse("nto_travel:landing") + '"')
        RegistrationPersonalData.objects.create(
            user=self.user, email=self.user.email, password=self.user.password,
            current_step="finish", nto_campaign_slug=self.campaign.slug,
        )
        # The return-to-season hint no longer determines program membership.
        response = self.client.get(reverse("index"))
        self.assertNotContains(response, 'href="' + reverse("nto_travel:landing") + '"')
        UserProgramPreferences.objects.create(user=self.user, nto_enabled=True, active_program='nto')
        response = self.client.get(reverse("index"))
        self.assertContains(response, 'href="' + reverse("nto_travel:landing") + '"')
        for main_url in ("apply", "candidate_testing_list", "motivation_letter", "my_video_page", "preparation"):
            self.assertNotContains(response, 'href="' + reverse(main_url) + '"')
        self.assertNotContains(response, 'id="selectionMenu"')
        self.assertNotContains(response, 'id="selectionDropdown"')
        self.assertNotContains(response, "Как проходит отбор")
        self.assertNotContains(response, "ещё не отправлена")
        info.form_status = "submitted"
        info.status = "FINAL STAGE"
        info.save(update_fields=["form_status", "status"])
        response = self.client.get(reverse("index"))
        self.assertContains(response, 'href="' + reverse("nto_travel:landing") + '"')
        self.assertNotContains(response, 'href="' + reverse("family_income:page") + '"')
        self.assertNotContains(response, 'href="' + reverse("documents_dashboard") + '"')

    def test_rendered_answers_and_submission_errors_are_visible(self):
        app = self.filled_application()
        self.assertEqual(FinanceForm(instance=app)["previous_final"].value(), "no")
        self.assertEqual(MotivationForm(instance=app)["refused_olympiads"].value(), "no")
        self.client.force_login(self.user)
        finance = self.client.get(reverse("nto_travel:apply_step", args=[self.campaign.slug, "finance"]))
        self.assertContains(finance, '<option value="no" selected>Нет</option>', html=True)
        app.nto_goal = ""
        app.save(update_fields=["nto_goal"])
        response = self.client.post(reverse("nto_travel:submit", args=[self.campaign.slug]), {"accept_consent": "on"}, follow=True)
        self.assertContains(response, "Цель участия в НТО: Обязательный ответ")
        self.assertContains(response, "Город финала")
        self.assertContains(response, "Томск")
        app.status = "clarification"
        app.clarification_due_at = timezone.now() + timedelta(days=1)
        app.public_comment = "Пожалуйста, уточните план поездки."
        app.save()
        response = self.client.get(reverse("nto_travel:apply_step", args=[self.campaign.slug, "personal"]))
        self.assertContains(response, app.public_comment)

    @patch("core.views.send_email_verification_code")
    def test_registration_remembers_nto_destination(self, mocked_email):
        url = reverse("register_initial") + "?nto=" + self.campaign.slug
        response = self.client.get(url)
        self.assertContains(response, f'name="nto" value="{self.campaign.slug}"')
        response = self.client.post(reverse("register_initial"), {
            "nto": self.campaign.slug,
            "email": "new-nto@example.test",
            "password": "StrongTestPassword2026!",
            "password_confirm": "StrongTestPassword2026!",
        })
        self.assertEqual(response.status_code, 302)
        user = get_user_model().objects.get(email="new-nto@example.test")
        self.assertEqual(user.registrationpersonaldata.nto_campaign_slug, self.campaign.slug)
        mocked_email.assert_called_once()
        user.registrationpersonaldata.current_step = "finish"
        user.registrationpersonaldata.save(update_fields=["current_step"])
        response = self.client.get(reverse("finish_registration"))
        self.assertRedirects(response, reverse('programs'))
        self.assertEqual(user.program_preferences.active_program, '')
        response = self.client.post(reverse('program_select'), {'program': 'nto'})
        self.assertRedirects(response, reverse("nto_travel:apply_step", args=[self.campaign.slug, "personal"]))
