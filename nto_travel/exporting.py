from datetime import datetime

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import NtoApplication
from .presentation import LABELS, answer_rows


REVIEW_COLUMNS = [
    ("review_criticality", "Критичность поддержки"), ("review_motivation", "Мотивация"),
    ("review_author", "Автор заключения"), ("readers", "Прочитали"),
]
REGISTRY = [
    ("pk", "Номер заявки"), ("full_name", "ФИО"), ("status", "Статус"),
    ("last_submitted_at", "Дата последней отправки"), ("region", "Субъект РФ"),
    ("city", "Город проживания"), ("school_class", "Класс"), ("first_profile", "Первый профиль"),
    ("final_city", "Город финала"), ("ticket_cost", "Стоимость билетов, ₽"),
    ("income_per_member", "Доход на человека, ₽/мес"),
    ("school_support", "Ответ школы"), ("parent_support", "Ответ родителей"),
    ("possible_payers", "Возможные источники оплаты"), ("clarification_due_at", "Срок уточнения"),
    ("decided_at", "Дата решения"),
] + REVIEW_COLUMNS
ANSWERS = [
    ("pk", "Номер заявки"), ("full_name", "ФИО"), ("status", "Статус"),
    ("profiles", "Профили по приоритету"),
] + [(key, label) for key, label in LABELS if key not in {"finalist_declared", "parent_contact_declared"}] + [
    ("public_comment", "Комментарий участнику"),
] + REVIEW_COLUMNS + [
    ("review_criticality_reason", "Обоснование критичности"),
    ("review_motivation_reason", "Обоснование мотивации"),
    ("review_updated_at", "Заключение изменено"),
]


def append_row(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"
        elif isinstance(cell.value, datetime):
            cell.number_format = "dd.mm.yyyy hh:mm"
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def application_workbook(applications):
    book = Workbook()
    book.remove(book.active)
    sheets = [(book.create_sheet("Реестр"), REGISTRY), (book.create_sheet("Ответы анкеты"), ANSWERS)]
    for sheet, columns in sheets:
        append_row(sheet, [label for _, label in columns])
    for app in applications:
        answers = dict(zip((key for key, _ in LABELS), (value for _, value in answer_rows(app, empty="", separator="\n"))))
        ranks = sorted(app.ranked_profiles.all(), key=lambda row: row.rank)
        use_snapshot = app.status != NtoApplication.Status.DRAFT
        names = [(row.name_snapshot or row.profile.name) if use_snapshot else row.profile.name for row in ranks]
        cities = [(row.city_snapshot or row.profile.final_city) if use_snapshot else row.profile.final_city for row in ranks]
        answers.update({
            "pk": app.pk, "full_name": " ".join(filter(None, [app.last_name, app.first_name, app.middle_name])),
            "status": app.get_status_display(), "first_profile": names[0] if names else "",
            "profiles": "\n".join(f"{row.rank}. {name} ({city})" for row, name, city in zip(ranks, names, cities)),
            "final_city": cities[0] if cities else app.final_city,
            "public_comment": app.public_comment,
        })
        review = getattr(app, "review", None)
        answers.update({
            "review_criticality": review.criticality if review else None,
            "review_motivation": review.motivation if review else None,
            "review_author": (review.author.get_full_name().strip() or review.author.get_username()) if review and review.author else "",
            "review_criticality_reason": review.criticality_reason if review else "",
            "review_motivation_reason": review.motivation_reason if review else "",
            "review_updated_at": timezone.localtime(review.updated_at).replace(tzinfo=None) if review else None,
            "readers": "\n".join(mark.user.get_full_name().strip() or mark.user.get_username() for mark in app.read_marks.all()),
        })
        for key in ("last_submitted_at", "clarification_due_at", "decided_at"):
            value = getattr(app, key)
            answers[key] = timezone.localtime(value).replace(tzinfo=None) if value else None
        for sheet, columns in sheets:
            append_row(sheet, [answers[key] for key, _ in columns])
    for index, (sheet, columns) in enumerate(sheets):
        sheet.freeze_panes = "C2"
        for column, (key, _) in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(column)].width = 16 if key == "pk" else 32
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24476A")
        sheet.row_dimensions[1].height = 45
        if sheet.max_row > 1:
            # Excel tables own their filters; a worksheet filter over the same
            # range makes Excel repair and remove the table on opening.
            table = Table(displayName=f"NtoApplications{index + 1}", ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
        else:
            sheet.auto_filter.ref = sheet.dimensions
    return book
